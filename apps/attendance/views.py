import json
from datetime import date, datetime, time
from io import BytesIO

from django.contrib.auth import authenticate, login
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from PIL import Image, UnidentifiedImageError

from apps.core.services import record_audit
from apps.employees.models import Employee
from apps.organization.models import OrganizationMembership
from apps.workforce.models import ShiftTemplate

from .models import AttendanceRecord
from .services import AttendanceCalculator

TIMEKEEPING_HEADERS = ('Employee Number', 'Attendance Date', 'Time In', 'Time Out', 'Status', 'Remarks')
MAX_CLOCK_PHOTO_BYTES = 5 * 1024 * 1024


def _can_import_timekeeping(request, membership):
    return request.user.is_superuser or membership.role in (OrganizationMembership.Role.SUPER_USER, OrganizationMembership.Role.HR)


def _validate_clock_photo(uploaded):
    if uploaded is None: raise ValueError('A camera photo is required before clocking.')
    if uploaded.size > MAX_CLOCK_PHOTO_BYTES: raise ValueError('The clock photo must be 5 MB or smaller.')
    if not getattr(uploaded, 'content_type', '').startswith('image/'): raise ValueError('The clock photo must be an image.')
    try:
        uploaded.seek(0); image = Image.open(uploaded); image.verify(); uploaded.seek(0)
    except (UnidentifiedImageError, OSError, ValueError): raise ValueError('The clock photo is not a valid image.')
    return uploaded


def _cell_datetime(attendance_date, value):
    if value in (None, ''): return None
    if isinstance(value, datetime): value = value.time()
    if isinstance(value, time): return timezone.make_aware(datetime.combine(attendance_date, value))
    if isinstance(value, str):
        try: parsed = datetime.fromisoformat(value.strip())
        except ValueError: parsed = datetime.combine(attendance_date, datetime.strptime(value.strip(), '%H:%M').time())
        return parsed if timezone.is_aware(parsed) else timezone.make_aware(parsed)
    raise ValueError('Time values must be Excel time cells or ISO date-time text.')


def _dashboard_sections(membership):
    sections = [('employee_directory','Employee Directory','View people, assignments, and organization structure.','view_employees'),('attendance','Attendance','Review and manage daily attendance records.','view_attendance'),('scheduling','Shift Scheduling','Create and assign shift schedules to employees.','manage_attendance'),('leave','Leave Approvals','Review and approve employee leave requests.','approve_leave'),('payroll','Payroll','Review payroll data and compensation operations.','view_payroll'),('reports','Reports','Open workforce and operational reports.','view_reports'),('onboarding','Onboarding','Track new-hire onboarding steps and completion status.','manage_employees'),('talent','Talent Lifecycle','Manage recruiting, performance, benefits, and offboarding.','manage_employees'),('operations','HR Operations','Manage documents, announcements, workflows, projects, and external connectors.','manage_employees'),('ess','Employee Self-Service','View your own profile, attendance, leave, and approved payroll.','view_self')]
    return [{'key': k, 'name': n, 'description': d} for k,n,d,p in sections if membership.has_permission(p)]


def _is_maintenance_user(request, membership): return request.user.is_superuser or membership.role == OrganizationMembership.Role.SUPER_USER

def dashboard_page(request):
    if not request.user.is_authenticated: return redirect('login')
    from apps.accounts.views import workspace_url_for_user
    destination = workspace_url_for_user(request.user)
    if destination != '/workspace/': return redirect(destination)
    return render(request, 'dashboard.html')

def attendance_page(request):
    if not request.user.is_authenticated: return redirect('login')
    from apps.accounts.views import workspace_url_for_user
    return render(request, 'attendance/attendance.html', {'workspace_url': workspace_url_for_user(request.user)})

def clock_page(request): return render(request, 'attendance/clock.html')


def _clock_state(employee):
    today = timezone.localdate()
    open_record = employee.attendance_records.filter(attendance_date=today, time_in__isnull=False, time_out__isnull=True).order_by('-time_in').first()
    latest_record = employee.attendance_records.filter(attendance_date=today).order_by('-created_at').first()
    last_action = 'CLOCKED_IN' if open_record else ('CLOCKED_OUT' if latest_record and latest_record.time_out else 'NOT_STARTED')
    last_action_at = open_record.time_in if open_record else (latest_record.time_out if latest_record else None)
    assignment = employee.assignments.filter(start_date__lte=today).filter(Q(end_date__isnull=True) | Q(end_date__gte=today)).order_by('-is_primary','-start_date').select_related('shift_template').first()
    effective_shift = latest_record.effective_shift if latest_record else (assignment.shift_template if assignment else None)
    cover_shifts = ShiftTemplate.objects.all().order_by('name')
    return {'employee_name':f'{employee.first_name} {employee.last_name}','employee_number':employee.employee_number,'last_action':last_action,'last_action_at':last_action_at.isoformat() if last_action_at else None,'attendance_date':today.isoformat(),'shift_name':effective_shift.name if effective_shift else None,'can_clock_in':open_record is None and employee.is_active,'can_clock_out':open_record is not None,'assignment_available':assignment is not None,'clock_in_mode':latest_record.clock_in_mode if latest_record else None,'clocking_note':'Scheduled shift' if assignment else 'No scheduled shift. You may still clock in for a cover or unscheduled shift; the punch will be flagged for schedule reconciliation before payroll approval.','clock_in_photo_captured':bool(open_record and open_record.clock_in_photo),'clock_out_photo_captured':bool(latest_record and latest_record.time_out and latest_record.clock_out_photo),'cover_shifts':[{'id':str(shift.id),'name':shift.name,'start_time':shift.start_time.strftime('%H:%M'),'end_time':shift.end_time.strftime('%H:%M')} for shift in cover_shifts]}


@require_http_methods(['POST'])
def clock_login(request):
    user = authenticate(request, username=request.POST.get('username','').strip(), password=request.POST.get('password',''))
    if user is None or not user.is_active: return JsonResponse({'detail':'Invalid username or password.'},status=401)
    try: employee=user.employee_profile
    except Employee.DoesNotExist: return JsonResponse({'detail':'This account is not linked to an employee profile.'},status=403)
    if not employee.is_active: return JsonResponse({'detail':'This employee profile is inactive.'},status=403)
    membership=employee.organization.organization_memberships.filter(user=user,is_active=True).first()
    if membership is None or not employee.organization.is_active: return JsonResponse({'detail':'This employee does not have active organization access.'},status=403)
    login(request,user); return JsonResponse(_clock_state(employee))


@require_http_methods(['GET','POST'])
def clock_action(request):
    if not request.user.is_authenticated: return JsonResponse({'detail':'Sign in before using the time clock.'},status=401)
    try: employee=request.user.employee_profile
    except Employee.DoesNotExist: return JsonResponse({'detail':'This account is not linked to an employee profile.'},status=403)
    if not employee.is_active or not employee.organization.is_active: return JsonResponse({'detail':'This employee is not eligible to record attendance.'},status=403)
    membership=employee.organization.organization_memberships.filter(user=request.user,is_active=True).first()
    if membership is None: return JsonResponse({'detail':'This employee does not have active organization access.'},status=403)
    if request.method=='GET': return JsonResponse(_clock_state(employee))
    action=request.POST.get('action')
    try: photo=_validate_clock_photo(request.FILES.get('photo'))
    except ValueError as error: return JsonResponse({'detail':str(error)},status=400)
    now=timezone.now(); attendance_date=timezone.localdate()
    if action=='CLOCK_IN':
        existing=employee.attendance_records.filter(attendance_date=attendance_date).first()
        if existing is not None:
            if existing.time_in and existing.time_out: return JsonResponse({'detail':'An attendance record already exists for today. Multiple shift segments are not enabled yet; ask HR to reconcile the record.'},status=409)
            if existing.time_in and existing.time_out is None: return JsonResponse({'detail':'You are already clocked in.'},status=409)
        assignment=employee.assignments.filter(start_date__lte=attendance_date).filter(Q(end_date__isnull=True)|Q(end_date__gte=attendance_date)).order_by('-is_primary','-start_date').select_related('shift_template').first()
        requested_mode=request.POST.get('clock_in_mode','').upper()
        clock_shift=None
        if requested_mode==AttendanceRecord.ClockInMode.COVER:
            shift_id=request.POST.get('shift_id')
            if shift_id:
                try: clock_shift=ShiftTemplate.objects.get(id=shift_id)
                except (ShiftTemplate.DoesNotExist, ValueError, TypeError): return JsonResponse({'detail':'Select a valid shift to cover.'},status=400)
            elif assignment is not None:
                clock_shift=assignment.shift_template
            mode=AttendanceRecord.ClockInMode.COVER
        elif assignment is None:
            if requested_mode != AttendanceRecord.ClockInMode.UNSCHEDULED: return JsonResponse({'detail':'No scheduled shift is assigned. Confirm this is a cover or unscheduled work shift and try again.'},status=400)
            mode=AttendanceRecord.ClockInMode.UNSCHEDULED
        else:
            mode=AttendanceRecord.ClockInMode.SCHEDULED
            clock_shift=assignment.shift_template
        record=AttendanceRecord.objects.create(employee=employee,assignment=assignment,clock_shift=clock_shift,clock_in_mode=mode,attendance_date=attendance_date,time_in=now,clock_in_photo=photo,status=AttendanceRecord.Status.PRESENT,remarks='' if mode==AttendanceRecord.ClockInMode.SCHEDULED else ('Cover shift clock-in; permanent assignment preserved. Payroll reconciliation required.' if mode==AttendanceRecord.ClockInMode.COVER else 'Unscheduled clock-in; schedule reconciliation required before payroll approval.'))
        record_audit(organization=employee.organization,actor=request.user,action='attendance.clocked_in',entity=record,details={'photo_required':True,'clock_in_mode':mode,'covered_shift_id':str(clock_shift.id) if clock_shift and mode==AttendanceRecord.ClockInMode.COVER else None})
        return JsonResponse(_clock_state(employee))
    if action=='CLOCK_OUT':
        record=employee.attendance_records.filter(time_in__isnull=False,time_out__isnull=True).order_by('-attendance_date','-time_in').first()
        if record is None: return JsonResponse({'detail':'You are not currently clocked in.'},status=409)
        record.time_out=now; record.clock_out_photo=photo; record.full_clean(); record.save(update_fields=('time_out','clock_out_photo','updated_at')); AttendanceCalculator.update_record(record)
        record_audit(organization=employee.organization,actor=request.user,action='attendance.clocked_out',entity=record,details={'photo_required':True,'clock_in_mode':record.clock_in_mode})
        return JsonResponse(_clock_state(employee))
    return JsonResponse({'detail':'Action must be CLOCK_IN or CLOCK_OUT.'},status=400)


def dashboard(request):
    if not request.user.is_authenticated: return JsonResponse({'detail':'Authentication credentials were not provided.'},status=401)
    membership=OrganizationMembership.objects.filter(user=request.user,is_active=True,organization__is_active=True).select_related('organization').first()
    if membership is None: return JsonResponse({'detail':'No active organization membership found.'},status=403)
    attendance_date=timezone.localdate(); today_records=AttendanceRecord.objects.filter(employee__organization=membership.organization,attendance_date=attendance_date).select_related('employee')
    return JsonResponse({'role':membership.role,'role_label':membership.get_role_display(),'is_maintenance_user':_is_maintenance_user(request,membership),'sections':_dashboard_sections(membership),'date':attendance_date.isoformat(),'active_employee_count':Employee.objects.filter(organization=membership.organization,is_active=True).count(),'today_attendance_count':today_records.count(),'today_present_count':today_records.filter(status=AttendanceRecord.Status.PRESENT).count(),'today_absent_count':today_records.filter(status=AttendanceRecord.Status.ABSENT).count(),'records':[{'employee_number':r.employee.employee_number,'employee_name':f'{r.employee.first_name} {r.employee.last_name}','status':r.status,'late_minutes':r.late_minutes,'undertime_minutes':r.undertime_minutes,'overtime_minutes':r.overtime_minutes} for r in today_records]})


@require_http_methods(['POST'])
def record_attendance(request):
    if not request.user.is_authenticated: return JsonResponse({'detail':'Authentication credentials were not provided.'},status=401)
    membership=OrganizationMembership.objects.filter(user=request.user,is_active=True,organization__is_active=True).select_related('organization').first()
    if membership is None or not membership.has_permission('manage_attendance'): return JsonResponse({'detail':'No active organization membership found.'},status=403)
    try:
        payload=json.loads(request.body); attendance_date=datetime.strptime(payload['attendance_date'],'%Y-%m-%d').date(); employee=Employee.objects.get(id=payload['employee_id'],organization=membership.organization,is_active=True)
    except (KeyError,TypeError,ValueError,json.JSONDecodeError): return JsonResponse({'detail':'Provide employee_id and attendance_date in valid format.'},status=400)
    except Employee.DoesNotExist: return JsonResponse({'detail':'Employee was not found in your organization.'},status=404)
    assignment=employee.assignments.filter(start_date__lte=attendance_date).filter(Q(end_date__isnull=True)|Q(end_date__gte=attendance_date)).order_by('-is_primary','-start_date').select_related('shift_template').first()
    if assignment is None: return JsonResponse({'detail':'Employee has no active assignment for this date.'},status=400)
    def parse_datetime(field_name):
        value=payload.get(field_name)
        if not value: return None
        parsed=datetime.fromisoformat(value); return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed
    try: time_in=parse_datetime('time_in'); time_out=parse_datetime('time_out')
    except (TypeError,ValueError): return JsonResponse({'detail':'time_in and time_out must be ISO datetime values.'},status=400)
    if time_in and timezone.localtime(time_in).date()!=attendance_date: return JsonResponse({'detail':'time_in must use the attendance date in Asia/Manila.'},status=400)
    if time_in and time_out and time_out<=time_in: return JsonResponse({'detail':'time_out must be later than time_in.'},status=400)
    record,_=AttendanceRecord.objects.update_or_create(employee=employee,attendance_date=attendance_date,defaults={'assignment':assignment,'clock_shift':assignment.shift_template,'clock_in_mode':AttendanceRecord.ClockInMode.SCHEDULED,'time_in':time_in,'time_out':time_out,'status':payload.get('status',AttendanceRecord.Status.PRESENT),'remarks':payload.get('remarks','')}); AttendanceCalculator.update_record(record); record_audit(organization=membership.organization,actor=request.user,action='attendance.recorded',entity=record,details={'employee':record.employee.employee_number,'attendance_date':attendance_date.isoformat()}); return JsonResponse({'id':str(record.id),'employee_id':str(employee.id),'attendance_date':attendance_date.isoformat(),'status':record.status,'late_minutes':record.late_minutes,'undertime_minutes':record.undertime_minutes,'overtime_minutes':record.overtime_minutes},status=201)


@require_http_methods(['GET'])
def timekeeping_template(request):
    if not request.user.is_authenticated: return JsonResponse({'detail':'Authentication credentials were not provided.'},status=401)
    membership=OrganizationMembership.objects.filter(user=request.user,is_active=True,organization__is_active=True).first()
    if membership is None or not _can_import_timekeeping(request,membership): return JsonResponse({'detail':'HR or Super User access is required.'},status=403)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    workbook=Workbook(); sheet=workbook.active; sheet.title='Timekeeping'; sheet.append(TIMEKEEPING_HEADERS); sheet.append(['EMP-000001',timezone.localdate(),time(8,0),time(17,0),'PRESENT','Optional note'])
    for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
    sheet['B2'].number_format='yyyy-mm-dd'; sheet['C2'].number_format=sheet['D2'].number_format='hh:mm'
    for column,width in zip('ABCDEF',(22,18,16,16,16,36)): sheet.column_dimensions[column].width=width
    buffer=BytesIO(); workbook.save(buffer); response=HttpResponse(buffer.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); response['Content-Disposition']='attachment; filename="bizflow-timekeeping-template.xlsx"'; return response


@require_http_methods(['POST'])
def import_timekeeping(request):
    if not request.user.is_authenticated: return JsonResponse({'detail':'Authentication credentials were not provided.'},status=401)
    membership=OrganizationMembership.objects.filter(user=request.user,is_active=True,organization__is_active=True).select_related('organization').first()
    if membership is None or not _can_import_timekeeping(request,membership): return JsonResponse({'detail':'HR or Super User access is required.'},status=403)
    uploaded=request.FILES.get('file')
    if uploaded is None or not uploaded.name.lower().endswith('.xlsx'): return JsonResponse({'detail':'Upload a .xlsx timekeeping file.'},status=400)
    if uploaded.size>5*1024*1024: return JsonResponse({'detail':'The workbook must be 5 MB or smaller.'},status=400)
    try:
        from openpyxl import load_workbook
        workbook=load_workbook(uploaded,read_only=True,data_only=True); sheet=workbook.active; headers=tuple(cell.value.strip() if isinstance(cell.value,str) else cell.value for cell in next(sheet.iter_rows(min_row=1,max_row=1)))
        if headers!=TIMEKEEPING_HEADERS: return JsonResponse({'detail':f'Headers must be exactly: {", ".join(TIMEKEEPING_HEADERS)}.'},status=400)
    except Exception: return JsonResponse({'detail':'The workbook could not be read. Use the supplied template.'},status=400)
    errors,prepared=[],[]; status_values={choice.value for choice in AttendanceRecord.Status}
    for row_number,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
        if not any(value not in (None,'') for value in row): continue
        try:
            employee_number,attendance_day,time_in,time_out,status,remarks=row
            if not isinstance(employee_number,str) or not employee_number.strip(): raise ValueError('Employee Number is required.')
            if isinstance(attendance_day,datetime): attendance_day=attendance_day.date()
            if not isinstance(attendance_day,date): raise ValueError('Attendance Date must be a date cell.')
            status=(status or AttendanceRecord.Status.PRESENT).strip().upper() if isinstance(status,str) else AttendanceRecord.Status.PRESENT
            if status not in status_values: raise ValueError('Status must be PRESENT, ABSENT, LEAVE, or HOLIDAY.')
            employee=Employee.objects.get(employee_number=employee_number.strip(),organization=membership.organization,is_active=True)
            assignment=employee.assignments.filter(start_date__lte=attendance_day).filter(Q(end_date__isnull=True)|Q(end_date__gte=attendance_day)).order_by('-is_primary','-start_date').select_related('shift_template').first()
            if assignment is None: raise ValueError('Employee has no active shift assignment on this date.')
            prepared.append((employee,assignment,attendance_day,_cell_datetime(attendance_day,time_in),_cell_datetime(attendance_day,time_out),status,str(remarks or '')))
        except Employee.DoesNotExist: errors.append({'row':row_number,'detail':'Employee Number was not found in your organization.'})
        except (TypeError,ValueError) as error: errors.append({'row':row_number,'detail':str(error)})
    if errors: return JsonResponse({'detail':'No records were imported. Fix the listed rows and upload again.','errors':errors},status=400)
    imported=0
    try:
        with transaction.atomic():
            for employee,assignment,attendance_day,time_in,time_out,status,remarks in prepared:
                if time_in and timezone.localtime(time_in).date()!=attendance_day: raise ValueError(f'{employee.employee_number}: Time In must match Attendance Date.')
                if time_in and time_out and time_out<=time_in: raise ValueError(f'{employee.employee_number}: Time Out must be later than Time In.')
                record,_=AttendanceRecord.objects.update_or_create(employee=employee,attendance_date=attendance_day,defaults={'assignment':assignment,'clock_shift':assignment.shift_template,'clock_in_mode':AttendanceRecord.ClockInMode.SCHEDULED,'time_in':time_in,'time_out':time_out,'status':status,'remarks':remarks}); record.full_clean(); AttendanceCalculator.update_record(record); imported+=1; record_audit(organization=membership.organization,actor=request.user,action='attendance.imported',entity=record,details={'source':uploaded.name,'row_count':len(prepared)})
    except ValueError as error: return JsonResponse({'detail':f'No records were imported: {error}'},status=400)
    return JsonResponse({'imported':imported,'updated_or_created':imported,'source':uploaded.name})
